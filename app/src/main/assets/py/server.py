# -*- coding: utf-8 -*-
"""Katri Locos POS - servidor local (solo biblioteca estándar + openpyxl para Excel)."""
import io
import json
import mimetypes
import os
import re
import socket
import sqlite3
import sys
import textwrap
import threading
from datetime import datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlparse

import impresora

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "datos")
STATIC = os.path.join(BASE, "static")
DB = os.path.join(DATA, "katri.db")
PORT = int(os.environ.get("KATRI_PORT", "8080"))
ANCHO_TICKET = 32  # papel de 58 mm


def _version():
    try:
        archivos = [os.path.abspath(__file__), os.path.join(BASE, "impresora.py"), os.path.join(STATIC, "app.js"),
                    os.path.join(STATIC, "offline.js")]
        return str(int(max(os.path.getmtime(f) for f in archivos if os.path.exists(f))))
    except Exception:
        return "0"


VERSION = _version()


# ----------------------------------------------------------------- base de datos
def db():
    c = sqlite3.connect(DB, timeout=15)
    c.row_factory = sqlite3.Row
    return c


_ctx = threading.local()
LOCAL = False       # True cuando corre dentro del celular (modo sin conexión)
USAR_WAL = True


def hora_actual():
    """Hora actual; al sincronizar pedidos hechos sin conexión se usa la hora en que se hicieron."""
    return getattr(_ctx, "reloj", None) or datetime.now()


def ahora():
    return hora_actual().strftime("%Y-%m-%d %H:%M:%S")


SCHEMA = """
create table if not exists categorias(
  id integer primary key autoincrement, nombre text not null, orden integer default 0);
create table if not exists productos(
  id integer primary key autoincrement, categoria_id integer, nombre text not null,
  descripcion text default '', precio integer not null default 0,
  config text default '{}', activo integer default 1, orden integer default 0,
  encargo integer default 0);
create table if not exists mesas(
  id integer primary key autoincrement, nombre text not null,
  ubicacion text default '', orden integer default 0);
create table if not exists cuentas(
  id integer primary key autoincrement, tipo text not null default 'mesa',
  mesa_id integer, cliente text default '', telefono text default '',
  direccion text default '', notas text default '', costo_domicilio integer default 0,
  bandeja integer default 0, estado text not null default 'abierta',
  mesero text default '', creada text, cerrada text, metodo_pago text default '',
  propina integer default 0, recibido integer default 0);
create table if not exists items(
  id integer primary key autoincrement, cuenta_id integer not null, producto_id integer,
  nombre text not null, cantidad integer not null default 1, precio_base integer not null,
  precio_unit integer not null, opciones text default '{}', ronda integer default 1,
  creado text);
create table if not exists config(clave text primary key, valor text);
create table if not exists empleados(
  id integer primary key autoincrement, nombre text not null,
  valor_hora integer not null default 0, activo integer default 1, orden integer default 0);
create table if not exists turnos(
  id integer primary key autoincrement, empleado_id integer not null,
  entrada text not null, salida text, valor_hora integer not null default 0);
create index if not exists ix_turnos_emp on turnos(empleado_id, entrada);
create table if not exists sync_ops(id text primary key, res text);
create table if not exists clientes(
  telefono text primary key, cliente text default '', direccion text default '',
  notas text default '', costo_domicilio integer default 0, actualizado text);
create index if not exists ix_items_cuenta on items(cuenta_id);
create index if not exists ix_cuentas_estado on cuentas(estado, cerrada);
"""

PROT3 = ["Carne de Res", "Pollo", "Birria"]


def G(nombre, opciones, cantidad=1):
    return {"nombre": nombre, "opciones": opciones, "cantidad": cantidad}


def menu_inicial():
    """Carta de Katri Locos Street Food."""
    q = lambda *x: list(x)
    return [
        ("Dorilocos", [
            ("Dorilocos Sencillos", 17000, "Lechuga, Sour Cream, Pico e Gallo, 1 Proteína",
             [G("Proteína", PROT3)], q("Lechuga", "Sour Cream", "Pico e Gallo"), 0),
            ("Dorilocos 1 Proteína", 20000,
             "Lechuga, Sour Cream, Pico e Gallo, 1 Proteína, Maicitos, Maduritos, Queso Mozzarella, Tajín",
             [G("Proteína", PROT3)],
             q("Lechuga", "Sour Cream", "Pico e Gallo", "Maicitos", "Maduritos", "Queso Mozzarella", "Tajín"), 0),
            ("Dorilocos 2 Proteínas", 24000,
             "Lechuga, Sour Cream, Mermelada de Tocineta, Pico e Gallo, 2 Proteínas, Maicitos, Maduritos, Queso Mozzarella, Tajín",
             [G("Proteína", PROT3, 2)],
             q("Lechuga", "Sour Cream", "Mermelada de Tocineta", "Pico e Gallo", "Maicitos", "Maduritos",
               "Queso Mozzarella", "Tajín"), 0),
            ("Dorilocos 2 Proteínas Extra Queso", 28000,
             "Lechuga, Sour Cream, Pico e Gallo, Mermelada de Tocineta, 2 Proteínas, Maicitos, Maduritos, Queso Mozzarella, Extra Queso, Tajín",
             [G("Proteína", PROT3, 2)],
             q("Lechuga", "Sour Cream", "Pico e Gallo", "Mermelada de Tocineta", "Maicitos", "Maduritos",
               "Queso Mozzarella", "Extra Queso", "Tajín"), 0),
            ("Dorilocos Familiar", 65000,
             "Lechuga, Sour Cream, Pico e Gallo, Mermelada de Tocineta, 2 Proteínas, Maicitos, Maduritos, Queso Mozzarella, Extra Queso, Tajín",
             [G("Proteína", PROT3, 2)],
             q("Lechuga", "Sour Cream", "Pico e Gallo", "Mermelada de Tocineta", "Maicitos", "Maduritos",
               "Queso Mozzarella", "Extra Queso", "Tajín"), 0),
        ]),
        ("Tacos", [
            ("Tacos x1", 10000, "Tortilla de harina o maíz, Lechuga, Sour Cream, Pico e Gallo, 1 Proteína, Queso Mozzarella",
             [G("Tortilla", ["Harina", "Maíz"]), G("Proteína", PROT3)],
             q("Lechuga", "Sour Cream", "Pico e Gallo", "Queso Mozzarella"), 0),
            ("Tacos x3", 29000, "Tortilla de harina o maíz, Lechuga, Sour Cream, Pico e Gallo, 1 Proteína por taco, Queso Mozzarella",
             [G("Tortilla", ["Harina", "Maíz"]), G("Proteína (una por taco)", PROT3, 3)],
             q("Lechuga", "Sour Cream", "Pico e Gallo", "Queso Mozzarella"), 0),
        ]),
        ("Maicitos", [
            ("Maicitos Pequeños", 17000, "Maicitos bañados en Queso Crema, Pollo, Queso Mozzarella, Ripio de Papa",
             [], q("Queso Crema", "Pollo", "Queso Mozzarella", "Ripio de Papa"), 0),
            ("Maicitos Grandes", 29000, "Maicitos bañados en Queso Crema, Pollo, Queso Mozzarella, Ripio de Papa",
             [], q("Queso Crema", "Pollo", "Queso Mozzarella", "Ripio de Papa"), 0),
        ]),
        ("Burrito y más", [
            ("Burrito", 20000, "Lechuga, Sour Cream, Pico e Gallo, Mermelada de Tocineta, 1 Proteína, Maicitos, Queso Mozzarella",
             [G("Proteína", PROT3)],
             q("Lechuga", "Sour Cream", "Pico e Gallo", "Mermelada de Tocineta", "Maicitos", "Queso Mozzarella"), 0),
            ("Lasagna", 30000, "Patacón maduro, Salsa Bechamel, Carne de Res y/o Pollo, Maicitos, Queso Mozzarella",
             [G("Proteína", ["Carne de Res", "Pollo", "Carne de Res y Pollo"])],
             q("Maicitos", "Queso Mozzarella"), 1),
            ("Patacones", 20000, "2 Patacones verdes o maduros, 1 Proteína, Guacamole, Pico e Gallo, Queso Mozzarella",
             [G("Patacón", ["Verdes", "Maduros"]), G("Proteína", PROT3)],
             q("Guacamole", "Pico e Gallo", "Queso Mozzarella"), 0),
        ]),
        ("Hamburguesas", [
            ("Hamburguesa", 26000,
             "Pan artesanal con base de Salsa de Leña, Lechuga, Pico e Gallo, 1 Proteína, Queso Mozzarella, Mermelada de Tocineta, Chicharroncitos artesanales, Ripio de Papa",
             [G("Proteína", PROT3 + ["Artesanal"])],
             q("Lechuga", "Pico e Gallo", "Queso Mozzarella", "Mermelada de Tocineta",
               "Chicharroncitos", "Ripio de Papa"), 0),
            ("Hamburguesa La Mera Mera", 28000,
             "Pan artesanal con base de Queso Crema, Lechuga, Pico e Gallo, Carne de Birria, Queso Mozzarella, Mermelada de Tocineta, Chicharroncitos, Ripio de Doritos",
             [], q("Lechuga", "Pico e Gallo", "Queso Mozzarella", "Mermelada de Tocineta",
                   "Chicharroncitos", "Ripio de Doritos"), 0),
            ("Hamburguesas Mini x1", 17000,
             "Pan artesanal con base de Queso Crema, 1 Proteína, Queso Mozzarella, Mermelada de Tocineta, Ripio de Doritos",
             [G("Proteína", PROT3)], q("Queso Mozzarella", "Mermelada de Tocineta", "Ripio de Doritos"), 0),
            ("Hamburguesas Mini x3", 46000,
             "3 minis con base de Queso Crema, 1 Proteína por mini, Queso Mozzarella, Mermelada de Tocineta, Ripio de Doritos",
             [G("Proteína (una por mini)", PROT3, 3)],
             q("Queso Mozzarella", "Mermelada de Tocineta", "Ripio de Doritos"), 0),
        ]),
        ("Salchipapas", [
            ("Salchipapa Porción", 8000, "", [], [], 0),
            ("Salchipapa Sencilla", 12000, "", [], [], 0),
            ("Salchipapa Mediana", 20000, "Maicitos, 1 Proteína, Queso Mozzarella",
             [G("Proteína", PROT3)], q("Maicitos", "Queso Mozzarella"), 0),
            ("Salchipapa Grande", 36000, "Maicitos, 1 Proteína, Queso Mozzarella",
             [G("Proteína", PROT3)], q("Maicitos", "Queso Mozzarella"), 0),
            ("Salchipapa Extra Familiar", 78000, "Maicitos, 1 Proteína, Queso Mozzarella",
             [G("Proteína", PROT3)], q("Maicitos", "Queso Mozzarella"), 0),
        ]),
        ("Adiciones", [
            ("Adición Queso Mozzarella", 4000, "", [], [], 0),
            ("Adición Carne de Res", 4000, "", [], [], 0),
            ("Adición Pollo", 4000, "", [], [], 0),
            ("Adición Birria", 4000, "", [], [], 0),
            ("Adición Mermelada de Tocineta", 3000, "", [], [], 0),
            ("Adición Maicitos", 2000, "", [], [], 0),
            ("Adición Maduritos", 2000, "", [], [], 0),
        ]),
        ("Bebidas", [
            ("Sodas Saborizadas", 10000, "", [], [], 0),
            ("Micheladas Saborizadas", 12000, "", [], [], 0),
            ("Jugo en Agua", 7000, "", [], [], 0),
            ("Jugo en Leche", 8000, "", [], [], 0),
            ("Milo Pequeño", 7000, "", [G("Servido", ["Caliente", "Frío"])], [], 0),
            ("Milo Grande", 10000, "", [G("Servido", ["Caliente", "Frío"])], [], 0),
            ("Coca-Cola", 5000, "", [], [], 0),
            ("Postobón 350 ml", 4000, "", [], [], 0),
            ("Soda Bretaña", 3000, "", [], [], 0),
            ("3 Cordilleras", 10000, "", [], [], 0),
            ("Cerveza Nacional", 6000, "", [], [], 0),
        ]),
        ("Licores", [
            ("Aguardiente Trago", 4000, "", [], [], 0),
            ("Aguardiente Media", 41000, "", [], [], 0),
            ("Ron Trago", 5000, "", [], [], 0),
            ("Ron Media", 45000, "", [], [], 0),
            ("Tequila Trago", 8000, "", [], [], 0),
            ("Tequila Botella", 91000, "", [], [], 0),
        ]),
        ("Postres", [
            ("Quesadilla de Arequipe", 14000, "Tortilla rellena de queso mozzarella con arequipe", [], [], 0),
            ("Brownie con Helado", 16000, "Brownie caliente con helado y decoración de arequipe", [], [], 0),
            ("Patacón (postre)", 12000, "1 patacón maduro con queso y bocadillo gratinado", [], [], 0),
        ]),
    ]


def init_db():
    os.makedirs(DATA, exist_ok=True)
    c = db()
    if USAR_WAL:
        c.execute("pragma journal_mode=wal")
    c.executescript(SCHEMA)
    if not c.execute("select 1 from config where clave='seed'").fetchone():
        for oc, (cat, prods) in enumerate(menu_inicial()):
            cur = c.execute("insert into categorias(nombre,orden) values(?,?)", (cat, oc))
            cid = cur.lastrowid
            for op, (nom, precio, desc, grupos, quitar, enc) in enumerate(prods):
                cfg = json.dumps({"grupos": grupos, "quitar": quitar}, ensure_ascii=False)
                c.execute(
                    "insert into productos(categoria_id,nombre,descripcion,precio,config,orden,encargo)"
                    " values(?,?,?,?,?,?,?)", (cid, nom, desc, precio, cfg, op, enc))
        for i in range(1, 19):
            c.execute("insert into mesas(nombre,ubicacion,orden) values(?,?,?)", ("Mesa %d" % i, "", i))
        defaults = {
            "seed": "1", "negocio": "Katri Locos Street Food", "telefono": "313 653 1708",
            "bandeja_domicilio": "1000",
        }
        for k, v in defaults.items():
            c.execute("insert or replace into config(clave,valor) values(?,?)", (k, v))
        c.commit()
    if not c.execute("select 1 from config where clave='seed_personal'").fetchone():
        for i, (nom, val) in enumerate([("Esposita", 10000), ("Marta", 10000), ("Mariana", 9000),
                                        ("Isabel", 9000), ("Dulce", 9000)], 1):
            c.execute("insert into empleados(nombre,valor_hora,orden) values(?,?,?)", (nom, val, i))
        c.execute("insert or replace into config(clave,valor) values('seed_personal','1')")
        c.commit()
    if not c.execute("select 1 from config where clave='mig_sabores'").fetchone():
        sodas = ["Cereza", "Tamarindo", "Manzana", "Mandarina", "Limón", "Mango", "Piña"]
        jugos = ["Tamarindo", "Mora", "Mango", "Mandarina", "Maracuyá", "Fresa", "Uva", "Piña", "Lulo", "Guanábana"]
        for nombre, sab in (("Sodas Saborizadas", sodas), ("Micheladas Saborizadas", sodas),
                            ("Jugo en Agua", jugos), ("Jugo en Leche", jugos)):
            p = c.execute("select id,config from productos where nombre=?", (nombre,)).fetchone()
            if not p:
                continue
            try:
                cfg = json.loads(p["config"] or "{}")
            except Exception:
                cfg = {}
            gr = cfg.get("grupos", [])
            if not any(g.get("nombre") == "Sabor" for g in gr):
                gr.append(G("Sabor", sab))
            cfg["grupos"] = gr
            c.execute("update productos set config=? where id=?", (json.dumps(cfg, ensure_ascii=False), p["id"]))
        c.execute("insert or replace into config(clave,valor) values('mig_sabores','1')")
        c.commit()
    if not c.execute("select 1 from config where clave='mig_salsas_tajin'").fetchone():
        salsas = ["Salsa Rosada", "Mayonesa", "Salsa de Ajo", "Salsa de Tomate", "Piña", "Picante", "Todas"]
        nuevos = {
            "Salchipapas": [G("Salsas", ["Con salsas", "Sin salsas"]),
                            dict(G("Cuáles salsas", salsas), multi=True, depende={"grupo": "Salsas", "valor": "Con salsas"}),
                            G("Tajín", ["Con Tajín", "Sin Tajín"])],
            "Dorilocos": [G("Tajín", ["Con Tajín", "Sin Tajín"])],
        }
        for cat, grupos_nuevos in nuevos.items():
            for p in c.execute("select p.id,p.config from productos p join categorias k on k.id=p.categoria_id where k.nombre=?", (cat,)).fetchall():
                try:
                    cfg = json.loads(p["config"] or "{}")
                except Exception:
                    cfg = {}
                gr = cfg.get("grupos", [])
                nombres = {g["nombre"] for g in gr}
                gr += [g for g in grupos_nuevos if g["nombre"] not in nombres]
                cfg["grupos"] = gr
                cfg["quitar"] = [x for x in cfg.get("quitar", []) if x.lower() not in ("tajín", "tajin")]
                c.execute("update productos set config=? where id=?", (json.dumps(cfg, ensure_ascii=False), p["id"]))
        c.execute("insert or replace into config(clave,valor) values('mig_salsas_tajin','1')")
        c.commit()
    c.close()


def cfg_get(c):
    return {r["clave"]: r["valor"] for r in c.execute("select clave,valor from config")}


# ----------------------------------------------------------------- lógica
class Err(Exception):
    pass


def dinero(n):
    return "$" + "{:,}".format(int(n)).replace(",", ".")


def prod_row(r):
    d = dict(r)
    try:
        cfg = json.loads(d.pop("config") or "{}")
    except Exception:
        cfg = {}
    d["grupos"] = cfg.get("grupos", [])
    d["quitar"] = cfg.get("quitar", [])
    return d


def get_menu(c, incluir_inactivos=False):
    out = []
    for cat in c.execute("select * from categorias order by orden,id"):
        q = "select * from productos where categoria_id=?"
        if not incluir_inactivos:
            q += " and activo=1"
        q += " order by orden,id"
        out.append({"id": cat["id"], "nombre": cat["nombre"],
                    "productos": [prod_row(p) for p in c.execute(q, (cat["id"],))]})
    return out


def items_cuenta(c, cuenta_id):
    res = []
    for r in c.execute("select * from items where cuenta_id=? order by ronda,id", (cuenta_id,)):
        d = dict(r)
        d["opciones"] = json.loads(d["opciones"] or "{}")
        d["total"] = d["precio_unit"] * d["cantidad"]
        res.append(d)
    return res


def cuenta_dict(c, r, con_items=True):
    d = dict(r)
    its = items_cuenta(c, d["id"])
    d["subtotal"] = sum(i["total"] for i in its)
    d["cargos"] = (d["bandeja"] + d["costo_domicilio"]) if d["tipo"] == "domicilio" else 0
    d["total"] = d["subtotal"] + d["cargos"]
    d["n_items"] = sum(i["cantidad"] for i in its)
    if con_items:
        d["items"] = its
    if d["mesa_id"]:
        m = c.execute("select nombre,ubicacion from mesas where id=?", (d["mesa_id"],)).fetchone()
        d["mesa"] = m["nombre"] if m else ""
        d["ubicacion"] = m["ubicacion"] if m else ""
    return d


def listar_mesas(c):
    abiertas = {}
    for r in c.execute("select * from cuentas where estado='abierta' and tipo='mesa'"):
        abiertas[r["mesa_id"]] = cuenta_dict(c, r, False)
    res = []
    for m in c.execute("select * from mesas order by orden,id"):
        d = dict(m)
        cu = abiertas.get(m["id"])
        d["cuenta"] = ({"id": cu["id"], "total": cu["total"], "n_items": cu["n_items"],
                        "creada": cu["creada"]} if cu else None)
        res.append(d)
    return res


def preparar_item(c, it):
    p = c.execute("select * from productos where id=?", (it.get("producto_id"),)).fetchone()
    if not p:
        raise Err("Producto no encontrado")
    pr = prod_row(p)
    cantidad = max(1, int(it.get("cantidad", 1)))
    extras = []
    for eid in it.get("extras", []) or []:
        e = c.execute("select id,nombre,precio from productos where id=?", (eid,)).fetchone()
        if e:
            extras.append({"id": e["id"], "nombre": e["nombre"], "precio": e["precio"]})
    opciones = {
        "grupos": {k: list(v) for k, v in (it.get("grupos") or {}).items() if v},
        "quitar": list(it.get("quitar") or []),
        "extras": extras,
        "nota": (it.get("nota") or "").strip(),
    }
    unit = pr["precio"] + sum(e["precio"] for e in extras)
    return pr["id"], pr["nombre"], cantidad, pr["precio"], unit, opciones


def agregar_items(c, cuenta_id, items):
    if not items:
        raise Err("El pedido está vacío")
    ronda = (c.execute("select coalesce(max(ronda),0) from items where cuenta_id=?", (cuenta_id,))
             .fetchone()[0]) + 1
    ts = ahora()
    for it in items:
        pid, nombre, cant, base, unit, op = preparar_item(c, it)
        c.execute("insert into items(cuenta_id,producto_id,nombre,cantidad,precio_base,precio_unit,opciones,ronda,creado)"
                  " values(?,?,?,?,?,?,?,?,?)",
                  (cuenta_id, pid, nombre, cant, base, unit, json.dumps(op, ensure_ascii=False), ronda, ts))
    return ronda


def texto_comanda(c, cuenta_id, ronda):
    cu = c.execute("select * from cuentas where id=?", (cuenta_id,)).fetchone()
    W = ANCHO_TICKET
    L = []
    L.append("COMANDA COCINA".center(W))
    L.append("=" * W)
    if cu["tipo"] == "mesa":
        m = c.execute("select * from mesas where id=?", (cu["mesa_id"],)).fetchone()
        L.append(m["nombre"].upper())
        if m["ubicacion"]:
            L.append(m["ubicacion"])
    else:
        L.append("DOMICILIO")
        if cu["cliente"]:
            L.append(cu["cliente"])
    if ronda > 1:
        L.append("*** ADICIONAL (pedido #%d) ***" % ronda)
    L.append(hora_actual().strftime("%d/%m/%Y %H:%M"))
    if cu["mesero"]:
        L.append("Atiende: " + cu["mesero"])
    L.append("-" * W)
    for i in c.execute("select * from items where cuenta_id=? and ronda=? order by id", (cuenta_id, ronda)):
        op = json.loads(i["opciones"] or "{}")
        L.extend(textwrap.wrap("%dx %s" % (i["cantidad"], i["nombre"].upper()), W,
                               subsequent_indent="   "))
        for g, vals in (op.get("grupos") or {}).items():
            cnt = {}
            for v in vals:
                cnt[v] = cnt.get(v, 0) + 1
            txt = ", ".join(("%dx %s" % (n, v)) if n > 1 else v for v, n in cnt.items())
            L.extend(textwrap.wrap("%s: %s" % (g.split(" (")[0], txt), W, initial_indent="  ",
                                   subsequent_indent="    "))
        for q in op.get("quitar") or []:
            L.extend(textwrap.wrap("SIN " + q.upper(), W, initial_indent="  ! ", subsequent_indent="    "))
        for e in op.get("extras") or []:
            L.extend(textwrap.wrap("+ " + e["nombre"].replace("Adición ", ""), W,
                                   initial_indent="  ", subsequent_indent="    "))
        if op.get("nota"):
            L.extend(textwrap.wrap("Nota: " + op["nota"], W, initial_indent="  ", subsequent_indent="    "))
        L.append("")
    if cu["tipo"] == "domicilio" and cu["notas"]:
        L.append("-" * W)
        L.extend(textwrap.wrap("Obs: " + cu["notas"], W))
    L.append("=" * W)
    return "\n".join(L)


def pedido_mesa(c, mesa_id, body):
    m = c.execute("select * from mesas where id=?", (mesa_id,)).fetchone()
    if not m:
        raise Err("Mesa no encontrada")
    cu = c.execute("select id from cuentas where mesa_id=? and tipo='mesa' and estado='abierta'",
                   (mesa_id,)).fetchone()
    if cu:
        cid = cu["id"]
        if body.get("mesero"):
            c.execute("update cuentas set mesero=? where id=? and mesero=''", (body["mesero"], cid))
    else:
        cid = c.execute("insert into cuentas(tipo,mesa_id,estado,mesero,creada) values('mesa',?,'abierta',?,?)",
                        (mesa_id, body.get("mesero", ""), ahora())).lastrowid
    ronda = agregar_items(c, cid, body.get("items"))
    c.commit()
    return respuesta_pedido(c, cid, ronda)


def respuesta_pedido(c, cid, ronda):
    cu = cuenta_dict(c, c.execute("select * from cuentas where id=?", (cid,)).fetchone(), False)
    texto = texto_comanda(c, cid, ronda)
    ids = [r[0] for r in c.execute("select id from items where cuenta_id=? and ronda=? order by id", (cid, ronda))]
    return {"cuenta_id": cid, "ronda": ronda, "total": cu["total"], "comanda": texto, "item_ids": ids,
            "impresion": impresora.imprimir(cfg_get(c), texto, doble_linea=2)}


def pedido_cuenta(c, cid, body):
    cu = c.execute("select * from cuentas where id=?", (cid,)).fetchone()
    if not cu or cu["estado"] != "abierta":
        raise Err("La cuenta no está abierta")
    ronda = agregar_items(c, cid, body.get("items"))
    c.commit()
    return respuesta_pedido(c, cid, ronda)


def nuevo_domicilio(c, body):
    cfg = cfg_get(c)
    if not (body.get("direccion") or "").strip():
        raise Err("Falta la dirección")
    cid = c.execute(
        "insert into cuentas(tipo,cliente,telefono,direccion,notas,costo_domicilio,bandeja,estado,mesero,creada)"
        " values('domicilio',?,?,?,?,?,?,'abierta',?,?)",
        ((body.get("cliente") or "").strip(), (body.get("telefono") or "").strip(),
         body["direccion"].strip(), (body.get("notas") or "").strip(),
         0, max(0, int(body.get("envase") if body.get("envase") not in (None, "") else (cfg.get("bandeja_domicilio", 0) or 0))),
         body.get("mesero", ""), ahora())).lastrowid
    ronda = agregar_items(c, cid, body.get("items"))
    guardar_cliente(c, body)
    c.commit()
    return respuesta_pedido(c, cid, ronda)


def solo_digitos(t):
    return re.sub(r"\D", "", str(t or ""))


def guardar_cliente(c, body):
    """Recuerda los datos del cliente por su teléfono para el próximo domicilio."""
    tel = solo_digitos(body.get("telefono"))
    if len(tel) < 7:
        return
    c.execute("insert or replace into clientes(telefono,cliente,direccion,notas,costo_domicilio,actualizado) values(?,?,?,?,?,?)",
              (tel, (body.get("cliente") or "").strip(), (body.get("direccion") or "").strip(),
               (body.get("notas") or "").strip(), int(body.get("costo_domicilio") or 0), ahora()))


def buscar_cliente(c, tel):
    r = c.execute("select * from clientes where telefono=?", (solo_digitos(tel),)).fetchone()
    return dict(r) if r else None


def cerrar_cuenta(c, cid, body):
    r = c.execute("select * from cuentas where id=?", (cid,)).fetchone()
    if not r or r["estado"] != "abierta":
        raise Err("La cuenta no está abierta")
    cu = cuenta_dict(c, r, False)
    if cu["n_items"] == 0:
        raise Err("La cuenta no tiene productos")
    metodo = (body.get("metodo_pago") or "").strip()
    if not metodo:
        raise Err("Elige el método de pago")
    propina = max(0, int(body.get("propina") or 0))
    recibido = max(0, int(body.get("recibido") or 0))
    c.execute("update cuentas set estado='cerrada', cerrada=?, metodo_pago=?, propina=?, recibido=? where id=?",
              (ahora(), metodo, propina, recibido, cid))
    c.commit()
    return {"ok": True, "total": cu["total"], "propina": propina}


# ----------------------------------------------------------------- reportes
def ventas_rango(c, desde, hasta):
    d0, d1 = desde + " 00:00:00", hasta + " 23:59:59"
    res = []
    for r in c.execute("select * from cuentas where estado='cerrada' and cerrada between ? and ? order by cerrada",
                       (d0, d1)):
        res.append(cuenta_dict(c, r, True))
    return res


def reporte(c, desde, hasta):
    cuentas = ventas_rango(c, desde, hasta)

    def acc(d, k, v, n=1, extra=0):
        a = d.setdefault(k, [0, 0, 0])
        a[0] += v
        a[1] += n
        a[2] += extra

    dia, prod, mesero, mesa, metodo, hora, tipo = {}, {}, {}, {}, {}, {}, {}
    ventas = propinas = items_n = cargos = 0
    for cu in cuentas:
        ventas += cu["total"]
        propinas += cu["propina"]
        cargos += cu["cargos"]
        f = cu["cerrada"][:10]
        acc(dia, f, cu["total"])
        acc(mesero, cu["mesero"] or "(sin nombre)", cu["total"])
        if cu["tipo"] == "mesa":
            etiqueta = cu.get("mesa", "Mesa") + (" · " + cu["ubicacion"] if cu.get("ubicacion") else "")
        else:
            etiqueta = "Domicilios"
        acc(mesa, etiqueta, cu["total"])
        acc(metodo, cu["metodo_pago"], cu["total"], 1, cu["propina"])
        acc(hora, cu["cerrada"][11:13] + ":00", cu["total"])
        acc(tipo, "Mesa" if cu["tipo"] == "mesa" else "Domicilio", cu["total"])
        for it in cu["items"]:
            items_n += it["cantidad"]
            acc(prod, it["nombre"], it["precio_base"] * it["cantidad"], it["cantidad"])
            for e in it["opciones"].get("extras", []):
                acc(prod, "(adición) " + e["nombre"].replace("Adición ", ""), e["precio"] * it["cantidad"],
                    it["cantidad"])
    n = len(cuentas)

    def lista(d, key="nombre", orden="valor"):
        out = [{key: k, "ventas": v[0], "cantidad": v[1], "extra": v[2]} for k, v in d.items()]
        out.sort(key=(lambda x: -x["ventas"]) if orden == "valor" else (lambda x: x[key]))
        return out

    return {
        "desde": desde, "hasta": hasta,
        "resumen": {"ventas": ventas, "cuentas": n, "ticket_promedio": (ventas // n if n else 0),
                    "propinas": propinas, "items": items_n, "cargos_domicilio": cargos},
        "por_dia": lista(dia, "fecha", "clave"),
        "por_producto": lista(prod, "nombre"),
        "por_mesero": lista(mesero, "nombre"),
        "por_mesa": lista(mesa, "nombre"),
        "por_metodo": lista(metodo, "nombre"),
        "por_hora": lista(hora, "hora", "clave"),
        "por_tipo": lista(tipo, "nombre"),
        "cuentas": [{"id": cu["id"], "cerrada": cu["cerrada"], "tipo": cu["tipo"],
                     "mesa": cu.get("mesa", "") if cu["tipo"] == "mesa" else "Domicilio: " + cu["cliente"],
                     "ubicacion": cu.get("ubicacion", ""), "mesero": cu["mesero"],
                     "items": cu["n_items"], "subtotal": cu["subtotal"], "cargos": cu["cargos"],
                     "total": cu["total"], "propina": cu["propina"], "metodo": cu["metodo_pago"]}
                    for cu in cuentas],
    }


def excel_bytes(desde, hasta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    c = db()
    try:
        r = reporte(c, desde, hasta)
        negocio = cfg_get(c).get("negocio", "Katri Locos")
    finally:
        c.close()
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="F6C233")
    head_font = Font(bold=True, color="111111")
    money = '"$"#,##0'

    def hoja(titulo, cols, filas, first=False, money_cols=()):
        ws = wb.active if first else wb.create_sheet()
        ws.title = titulo
        ws.append(cols)
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center")
        for f in filas:
            ws.append(f)
        for col in money_cols:
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                for cell in row:
                    cell.number_format = money
        for i, col in enumerate(ws.columns, 1):
            ancho = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, ancho + 3), 48)
        ws.freeze_panes = "A2"
        return ws

    s = r["resumen"]
    ws = hoja("Resumen", ["Concepto", "Valor"], [
        ["Negocio", negocio], ["Desde", desde], ["Hasta", hasta],
        ["Ventas (sin propina)", s["ventas"]], ["Cuentas cerradas", s["cuentas"]],
        ["Ticket promedio", s["ticket_promedio"]], ["Productos vendidos", s["items"]],
        ["Cargos de domicilio incluidos", s["cargos_domicilio"]], ["Propinas", s["propinas"]],
    ], first=True)
    for fila in (5, 7, 9, 10):
        ws.cell(row=fila, column=2).number_format = money
    hoja("Por día", ["Fecha", "Cuentas", "Ventas"],
         [[x["fecha"], x["cantidad"], x["ventas"]] for x in r["por_dia"]], money_cols=(3,))
    hoja("Por producto", ["Producto", "Cantidad", "Ventas"],
         [[x["nombre"], x["cantidad"], x["ventas"]] for x in r["por_producto"]], money_cols=(3,))
    hoja("Por mesero", ["Mesero", "Cuentas", "Ventas"],
         [[x["nombre"], x["cantidad"], x["ventas"]] for x in r["por_mesero"]], money_cols=(3,))
    hoja("Por mesa", ["Mesa / ubicación", "Cuentas", "Ventas"],
         [[x["nombre"], x["cantidad"], x["ventas"]] for x in r["por_mesa"]], money_cols=(3,))
    hoja("Por método de pago", ["Método", "Cuentas", "Ventas", "Propinas"],
         [[x["nombre"], x["cantidad"], x["ventas"], x["extra"]] for x in r["por_metodo"]], money_cols=(3, 4))
    hoja("Por hora", ["Hora", "Cuentas", "Ventas"],
         [[x["hora"], x["cantidad"], x["ventas"]] for x in r["por_hora"]], money_cols=(3,))
    hoja("Detalle de cuentas",
         ["N°", "Cierre", "Mesa / domicilio", "Ubicación", "Mesero", "Productos", "Subtotal", "Cargos",
          "Total", "Propina", "Método"],
         [[x["id"], x["cerrada"], x["mesa"], x["ubicacion"], x["mesero"], x["items"], x["subtotal"],
           x["cargos"], x["total"], x["propina"], x["metodo"]] for x in r["cuentas"]],
         money_cols=(7, 8, 9, 10))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ----------------------------------------------------------------- personal y horas
FMT = "%Y-%m-%d %H:%M:%S"


def _dt(s):
    return datetime.strptime(s, FMT)


def segundos_turno(t, ahora_dt=None):
    fin = _dt(t["salida"]) if t["salida"] else (ahora_dt or hora_actual())
    # minutos completos: el pago no lleva segundos sueltos
    return max(0, int((fin - _dt(t["entrada"])).total_seconds() // 60)) * 60


def pago_turno(t, ahora_dt=None):
    return int(round(segundos_turno(t, ahora_dt) * t["valor_hora"] / 3600.0))


def dur_txt(seg):
    m = int(seg // 60)
    return "%d h %02d min" % (m // 60, m % 60)


def listar_personal(c):
    ahora_dt = hora_actual()
    hoy = ahora_dt.strftime("%Y-%m-%d")
    res = []
    for e in c.execute("select * from empleados where activo=1 order by orden,id"):
        d = dict(e)
        ab = c.execute("select * from turnos where empleado_id=? and salida is null order by id desc limit 1",
                       (e["id"],)).fetchone()
        d["turno_abierto"] = None
        if ab:
            d["turno_abierto"] = {"id": ab["id"], "entrada": ab["entrada"], "segundos": segundos_turno(ab, ahora_dt),
                                  "pago": pago_turno(ab, ahora_dt)}
        seg = pago = 0
        for t in c.execute("select * from turnos where empleado_id=? and substr(entrada,1,10)=?", (e["id"], hoy)):
            seg += segundos_turno(t, ahora_dt)
            pago += pago_turno(t, ahora_dt)
        d["hoy"] = {"segundos": seg, "pago": pago}
        res.append(d)
    return res


def guardar_empleado(c, b):
    nombre = (b.get("nombre") or "").strip()
    if not nombre:
        raise Err("Falta el nombre")
    try:
        valor = int(round(float(b.get("valor_hora") or 0)))
    except (TypeError, ValueError):
        raise Err("El valor de la hora no es válido")
    if valor < 0:
        raise Err("El valor de la hora no puede ser negativo")
    if b.get("id"):
        c.execute("update empleados set nombre=?, valor_hora=? where id=?", (nombre, valor, int(b["id"])))
        eid = int(b["id"])
    else:
        orden = c.execute("select coalesce(max(orden),0)+1 from empleados").fetchone()[0]
        eid = c.execute("insert into empleados(nombre,valor_hora,orden) values(?,?,?)", (nombre, valor, orden)).lastrowid
    c.commit()
    return {"id": eid}


def quitar_empleado(c, eid):
    if c.execute("select 1 from turnos where empleado_id=? and salida is null", (eid,)).fetchone():
        raise Err("Tiene un turno abierto: registra primero su salida")
    if c.execute("select 1 from turnos where empleado_id=?", (eid,)).fetchone():
        c.execute("update empleados set activo=0 where id=?", (eid,))  # se conserva el historial de pagos
    else:
        c.execute("delete from empleados where id=?", (eid,))
    c.commit()
    return {"ok": True}


def marcar_entrada(c, eid):
    e = c.execute("select * from empleados where id=? and activo=1", (eid,)).fetchone()
    if not e:
        raise Err("Empleado no encontrado")
    if c.execute("select 1 from turnos where empleado_id=? and salida is null", (eid,)).fetchone():
        raise Err("%s ya tiene el turno iniciado" % e["nombre"])
    ts = ahora()
    c.execute("insert into turnos(empleado_id,entrada,valor_hora) values(?,?,?)", (eid, ts, e["valor_hora"]))
    c.commit()
    return {"nombre": e["nombre"], "entrada": ts, "turno_id": c.execute("select max(id) from turnos where empleado_id=?", (eid,)).fetchone()[0]}


def marcar_salida(c, eid):
    e = c.execute("select * from empleados where id=?", (eid,)).fetchone()
    t = c.execute("select * from turnos where empleado_id=? and salida is null order by id desc limit 1",
                  (eid,)).fetchone()
    if not e or not t:
        raise Err("No tiene un turno iniciado")
    ts = ahora()
    c.execute("update turnos set salida=? where id=?", (ts, t["id"]))
    c.commit()
    t = c.execute("select * from turnos where id=?", (t["id"],)).fetchone()
    return {"nombre": e["nombre"], "entrada": t["entrada"], "salida": ts,
            "segundos": segundos_turno(t), "pago": pago_turno(t)}


def reporte_turnos(c, desde, hasta):
    ahora_dt = hora_actual()
    d0, d1 = desde + " 00:00:00", hasta + " 23:59:59"
    por_emp, turnos = {}, []
    for t in c.execute("select t.*, e.nombre from turnos t join empleados e on e.id=t.empleado_id "
                       "where t.entrada between ? and ? order by t.entrada", (d0, d1)):
        seg, pago = segundos_turno(t, ahora_dt), pago_turno(t, ahora_dt)
        a = por_emp.setdefault(t["empleado_id"], {"id": t["empleado_id"], "nombre": t["nombre"], "turnos": 0,
                                                  "segundos": 0, "pago": 0, "valor_hora": t["valor_hora"]})
        a["turnos"] += 1
        a["segundos"] += seg
        a["pago"] += pago
        a["valor_hora"] = t["valor_hora"]
        turnos.append({"id": t["id"], "empleado_id": t["empleado_id"], "empleado": t["nombre"],
                       "entrada": t["entrada"], "salida": t["salida"], "abierto": t["salida"] is None,
                       "segundos": seg, "pago": pago, "valor_hora": t["valor_hora"]})
    emps = sorted(por_emp.values(), key=lambda x: x["nombre"].lower())
    return {"desde": desde, "hasta": hasta, "empleados": emps, "turnos": turnos,
            "total_segundos": sum(x["segundos"] for x in emps), "total_pago": sum(x["pago"] for x in emps)}


def _fecha_hora(txt):
    txt = (txt or "").strip().replace("T", " ")
    if len(txt) == 16:
        txt += ":00"
    try:
        return _dt(txt).strftime(FMT)
    except ValueError:
        raise Err("Fecha u hora no válida")


def editar_turno(c, tid, b):
    t = c.execute("select * from turnos where id=?", (tid,)).fetchone()
    if not t:
        raise Err("Turno no encontrado")
    entrada = _fecha_hora(b.get("entrada"))
    salida = _fecha_hora(b.get("salida")) if (b.get("salida") or "").strip() else None
    if salida and salida < entrada:
        raise Err("La salida no puede ser antes de la entrada")
    if salida is None and c.execute("select 1 from turnos where empleado_id=? and salida is null and id<>?",
                                    (t["empleado_id"], tid)).fetchone():
        raise Err("Ya hay otro turno abierto de esta persona")
    c.execute("update turnos set entrada=?, salida=? where id=?", (entrada, salida, tid))
    c.commit()
    return {"ok": True}


def excel_personal_bytes(desde, hasta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    c = db()
    try:
        r = reporte_turnos(c, desde, hasta)
    finally:
        c.close()
    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="F6C233"), Font(bold=True, color="111111")
    money = '"$"#,##0'

    def hoja(titulo, cols, filas, first=False, money_cols=()):
        ws = wb.active if first else wb.create_sheet()
        ws.title = titulo
        ws.append(cols)
        for cell in ws[1]:
            cell.fill, cell.font, cell.alignment = head_fill, head_font, Alignment(horizontal="center")
        for f in filas:
            ws.append(f)
        for col in money_cols:
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                for cell in row:
                    cell.number_format = money
        for col in ws.columns:
            ancho = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, ancho + 3), 40)
        ws.freeze_panes = "A2"
        return ws

    filas = [[x["nombre"], x["turnos"], round(x["segundos"] / 3600.0, 2), dur_txt(x["segundos"]),
              x["valor_hora"], x["pago"]] for x in r["empleados"]]
    filas.append(["TOTAL", sum(x["turnos"] for x in r["empleados"]), round(r["total_segundos"] / 3600.0, 2),
                  dur_txt(r["total_segundos"]), "", r["total_pago"]])
    ws = hoja("Pago por persona", ["Persona", "Turnos", "Horas", "Tiempo", "Valor hora", "A pagar"], filas,
              first=True, money_cols=(5, 6))
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    hoja("Detalle de turnos", ["Persona", "Entrada", "Salida", "Tiempo", "Horas", "Valor hora", "Pago"],
         [[t["empleado"], t["entrada"], t["salida"] or "(en curso)", dur_txt(t["segundos"]),
           round(t["segundos"] / 3600.0, 2), t["valor_hora"], t["pago"]] for t in r["turnos"]], money_cols=(6, 7))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ----------------------------------------------------------------- IP local
def ip_local():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    except Exception:
        return "127.0.0.1"
    try:
        s.connect(("10.255.255.255", 1))
        return s.getsockname()[0]
    except Exception:
        return "127.0.0.1"
    finally:
        s.close()



def despachar(metodo, path, q, body):
    """Ejecuta una ruta de la API. Devuelve (estado_http, respuesta). Lo usa el servidor y el modo sin conexión."""
    for m, rx, fn in RUTAS:
        if m != metodo:
            continue
        mt = rx.match(path)
        if not mt:
            continue
        c = db()
        try:
            return 200, fn(c, mt.groupdict(), body, q)
        except Err as e:
            return 400, {"error": str(e)}
        except Exception as e:  # noqa
            return 500, {"error": "Error interno: %s" % e}
        finally:
            c.close()
    return 404, {"error": "No encontrado"}


# ----------------------------------------------------------------- copia y sincronización (celular sin conexión)
TABLAS = ["categorias", "productos", "mesas", "cuentas", "items", "config", "empleados", "turnos", "clientes"]


def snapshot(c):
    out = {}
    for t in TABLAS:
        out[t] = [dict(r) for r in c.execute("select * from %s" % t)]
    return out


def restaurar(c, snap):
    """Reemplaza TODOS los datos locales por la copia del PC (solo se permite dentro del celular)."""
    if not LOCAL:
        raise Err("Operación no permitida")
    for t in TABLAS:
        c.execute("delete from %s" % t)
    for t in TABLAS:
        filas = snap.get(t) or []
        if not filas:
            continue
        cols = list(filas[0].keys())
        c.executemany("insert into %s(%s) values(%s)" % (t, ",".join(cols), ",".join("?" * len(cols))),
                      [tuple(f.get(k) for k in cols) for f in filas])
    # lo creado sin conexión usa números altos para no chocar con los del PC
    for t in ("categorias", "productos", "mesas", "cuentas", "items", "empleados", "turnos"):
        mx = c.execute("select coalesce(max(id),0) from %s" % t).fetchone()[0]
        seq = max(1000000, mx)
        if c.execute("select 1 from sqlite_sequence where name=?", (t,)).fetchone():
            c.execute("update sqlite_sequence set seq=? where name=?", (seq, t))
        else:
            c.execute("insert into sqlite_sequence(name,seq) values(?,?)", (t, seq))
    c.commit()
    return {"ok": True}


ENTIDAD_RUTA = {"mesas": "mesas", "cuentas": "cuentas", "items": "items", "productos": "productos",
                "categorias": "categorias", "personal": "empleados", "turnos": "turnos"}


def _mapear(op, mapa):
    ruta = op["path"]
    partes = ruta.split("/")  # ['', 'api', entidad, id, ...]
    if len(partes) > 3 and partes[2] in ENTIDAD_RUTA and partes[3].isdigit():
        real = mapa.get(ENTIDAD_RUTA[partes[2]], {}).get(int(partes[3]))
        if real:
            partes[3] = str(real)
    cuerpo = json.loads(json.dumps(op.get("body") or {}))
    for k, ent in (("producto_id", "productos"), ("categoria_id", "categorias"), ("mesa_id", "mesas")):
        if isinstance(cuerpo.get(k), int) and mapa.get(ent, {}).get(cuerpo[k]):
            cuerpo[k] = mapa[ent][cuerpo[k]]
    if isinstance(cuerpo.get("extras"), list):
        cuerpo["extras"] = [mapa.get("productos", {}).get(x, x) for x in cuerpo["extras"]]
    if cuerpo.get("id") and isinstance(cuerpo["id"], int) and partes[2:3] and partes[2] in ENTIDAD_RUTA:
        cuerpo["id"] = mapa.get(ENTIDAD_RUTA[partes[2]], {}).get(cuerpo["id"], cuerpo["id"])
    return "/".join(partes), cuerpo


def _registrar_ids(op, local, real, mapa):
    """Relaciona los números creados en el celular con los reales del PC."""
    if not isinstance(local, dict) or not isinstance(real, dict):
        return
    if local.get("cuenta_id") and real.get("cuenta_id"):
        mapa.setdefault("cuentas", {})[local["cuenta_id"]] = real["cuenta_id"]
    for a, b in zip(local.get("item_ids") or [], real.get("item_ids") or []):
        mapa.setdefault("items", {})[a] = b
    partes = op["path"].split("/")
    if local.get("turno_id") and real.get("turno_id"):
        mapa.setdefault("turnos", {})[local["turno_id"]] = real["turno_id"]
    if op["metodo"] == "POST" and len(partes) == 3 and partes[2] in ENTIDAD_RUTA and local.get("id") and real.get("id"):
        mapa.setdefault(ENTIDAD_RUTA[partes[2]], {})[local["id"]] = real["id"]


def sincronizar(ops):
    """Aplica en el PC, en orden, lo que se hizo en el celular sin conexión."""
    mapa, resultados, errores = {}, [], []
    ahora_real = datetime.now()
    for op in ops:
        # si esta operación ya se aplicó (la respuesta anterior se perdió), no se repite
        c0 = db()
        try:
            previa = c0.execute("select res from sync_ops where id=?", (str(op.get("id")),)).fetchone()
        finally:
            c0.close()
        if previa:
            _registrar_ids(op, op.get("local"), json.loads(previa["res"]), mapa)
            resultados.append({"id": op.get("id"), "ok": True, "repetida": True})
            continue
        ruta, cuerpo = _mapear(op, mapa)
        try:
            t = datetime.strptime(op["ts"], "%Y-%m-%d %H:%M:%S")
        except Exception:
            t = None
        _ctx.reloj = t
        impresora._ctx.omitir = bool(op.get("impreso") or (t and (ahora_real - t).total_seconds() > 600))
        try:
            estado, res = despachar(op["metodo"], ruta, op.get("query") or {}, cuerpo)
        finally:
            _ctx.reloj = None
            impresora._ctx.omitir = False
        if estado == 200:
            _registrar_ids(op, op.get("local"), res, mapa)
            c1 = db()
            try:
                c1.execute("insert or replace into sync_ops(id,res) values(?,?)", (str(op.get("id")), json.dumps(res, ensure_ascii=False)))
                c1.commit()
            finally:
                c1.close()
            resultados.append({"id": op.get("id"), "ok": True})
        else:
            err = (res or {}).get("error", "Error")
            errores.append({"id": op.get("id"), "resumen": op.get("resumen") or op["path"], "error": err})
            resultados.append({"id": op.get("id"), "ok": False, "error": err})
    return {"resultados": resultados, "errores": errores, "aplicadas": sum(1 for r in resultados if r["ok"])}


# ----------------------------------------------------------------- HTTP
RUTAS = []


def ruta(metodo, patron):
    def deco(fn):
        RUTAS.append((metodo, re.compile("^" + patron + "$"), fn))
        return fn
    return deco


@ruta("GET", "/api/info")
def r_info(c, p, b, q):
    cfg = cfg_get(c)
    ip = ip_local()
    return {"ip": ip, "puerto": PORT, "url_mesero": "http://%s:%d/mesero" % (ip, PORT),
            "negocio": cfg.get("negocio", ""), "telefono": cfg.get("telefono", ""),
            "bandeja_domicilio": int(cfg.get("bandeja_domicilio", 0) or 0)}


@ruta("GET", "/api/menu")
def r_menu(c, p, b, q):
    return get_menu(c, incluir_inactivos=q.get("todos") == "1")


@ruta("GET", "/api/mesas")
def r_mesas(c, p, b, q):
    return listar_mesas(c)


@ruta("POST", "/api/mesas/(?P<id>\\d+)/pedido")
def r_pedido_mesa(c, p, b, q):
    return pedido_mesa(c, int(p["id"]), b)


@ruta("GET", "/api/cuentas")
def r_cuentas(c, p, b, q):
    sql, args = "select * from cuentas where 1=1", []
    if q.get("estado"):
        sql += " and estado=?"
        args.append(q["estado"])
    if q.get("tipo"):
        sql += " and tipo=?"
        args.append(q["tipo"])
    if q.get("hoy") == "1":
        sql += " and (estado='abierta' or substr(cerrada,1,10)=?)"
        args.append(hora_actual().strftime("%Y-%m-%d"))
    sql += " order by id desc"
    return [cuenta_dict(c, r, False) for r in c.execute(sql, args)]


@ruta("GET", "/api/cuentas/(?P<id>\\d+)")
def r_cuenta(c, p, b, q):
    r = c.execute("select * from cuentas where id=?", (int(p["id"]),)).fetchone()
    if not r:
        raise Err("Cuenta no encontrada")
    return cuenta_dict(c, r)


@ruta("POST", "/api/cuentas/(?P<id>\\d+)/pedido")
def r_pedido_cuenta(c, p, b, q):
    return pedido_cuenta(c, int(p["id"]), b)


@ruta("POST", "/api/cuentas/(?P<id>\\d+)/cerrar")
def r_cerrar(c, p, b, q):
    return cerrar_cuenta(c, int(p["id"]), b)


@ruta("POST", "/api/cuentas/(?P<id>\\d+)/anular")
def r_anular(c, p, b, q):
    c.execute("update cuentas set estado='anulada', cerrada=? where id=? and estado='abierta'",
              (ahora(), int(p["id"])))
    c.commit()
    return {"ok": True}


@ruta("DELETE", "/api/items/(?P<id>\\d+)")
def r_del_item(c, p, b, q):
    it = c.execute("select i.id from items i join cuentas u on u.id=i.cuenta_id "
                   "where i.id=? and u.estado='abierta'", (int(p["id"]),)).fetchone()
    if not it:
        raise Err("No se puede quitar ese producto")
    c.execute("delete from items where id=?", (it["id"],))
    c.commit()
    return {"ok": True}


@ruta("POST", "/api/domicilios")
def r_domicilio(c, p, b, q):
    return nuevo_domicilio(c, b)


@ruta("GET", "/api/reportes")
def r_reporte(c, p, b, q):
    hoy = hora_actual().strftime("%Y-%m-%d")
    return reporte(c, q.get("desde") or hoy, q.get("hasta") or hoy)


@ruta("POST", "/api/productos")
def r_producto(c, p, b, q):
    nombre = (b.get("nombre") or "").strip()
    if not nombre:
        raise Err("Falta el nombre")
    cfg = json.dumps({"grupos": b.get("grupos") or [], "quitar": b.get("quitar") or []}, ensure_ascii=False)
    vals = (b.get("categoria_id"), nombre, (b.get("descripcion") or "").strip(), int(b.get("precio") or 0),
            cfg, 1 if b.get("activo", 1) else 0, 1 if b.get("encargo") else 0)
    if b.get("id"):
        c.execute("update productos set categoria_id=?,nombre=?,descripcion=?,precio=?,config=?,activo=?,encargo=?"
                  " where id=?", vals + (int(b["id"]),))
        pid = int(b["id"])
    else:
        orden = c.execute("select coalesce(max(orden),0)+1 from productos where categoria_id=?",
                          (b.get("categoria_id"),)).fetchone()[0]
        pid = c.execute("insert into productos(categoria_id,nombre,descripcion,precio,config,activo,encargo,orden)"
                        " values(?,?,?,?,?,?,?,?)", vals + (orden,)).lastrowid
    c.commit()
    return {"id": pid}


@ruta("DELETE", "/api/productos/(?P<id>\\d+)")
def r_del_producto(c, p, b, q):
    c.execute("delete from productos where id=?", (int(p["id"]),))
    c.commit()
    return {"ok": True}


@ruta("POST", "/api/categorias")
def r_categoria(c, p, b, q):
    nombre = (b.get("nombre") or "").strip()
    if not nombre:
        raise Err("Falta el nombre")
    if b.get("id"):
        c.execute("update categorias set nombre=? where id=?", (nombre, int(b["id"])))
        cid = int(b["id"])
    else:
        orden = c.execute("select coalesce(max(orden),0)+1 from categorias").fetchone()[0]
        cid = c.execute("insert into categorias(nombre,orden) values(?,?)", (nombre, orden)).lastrowid
    c.commit()
    return {"id": cid}


@ruta("DELETE", "/api/categorias/(?P<id>\\d+)")
def r_del_categoria(c, p, b, q):
    if c.execute("select 1 from productos where categoria_id=?", (int(p["id"]),)).fetchone():
        raise Err("La categoría tiene productos; muévelos o elimínalos primero")
    c.execute("delete from categorias where id=?", (int(p["id"]),))
    c.commit()
    return {"ok": True}


@ruta("POST", "/api/mesas")
def r_mesa(c, p, b, q):
    nombre = (b.get("nombre") or "").strip()
    if not nombre:
        raise Err("Falta el nombre de la mesa")
    ub = (b.get("ubicacion") or "").strip()
    if b.get("id"):
        c.execute("update mesas set nombre=?, ubicacion=? where id=?", (nombre, ub, int(b["id"])))
        mid = int(b["id"])
    else:
        orden = c.execute("select coalesce(max(orden),0)+1 from mesas").fetchone()[0]
        mid = c.execute("insert into mesas(nombre,ubicacion,orden) values(?,?,?)", (nombre, ub, orden)).lastrowid
    c.commit()
    return {"id": mid}


@ruta("DELETE", "/api/mesas/(?P<id>\\d+)")
def r_del_mesa(c, p, b, q):
    if c.execute("select 1 from cuentas where mesa_id=? and estado='abierta'", (int(p["id"]),)).fetchone():
        raise Err("La mesa tiene una cuenta abierta")
    c.execute("delete from mesas where id=?", (int(p["id"]),))
    c.commit()
    return {"ok": True}


@ruta("GET", "/api/personal")
def r_personal(c, p, b, q):
    return listar_personal(c)


@ruta("POST", "/api/personal")
def r_personal_guardar(c, p, b, q):
    return guardar_empleado(c, b)


@ruta("DELETE", "/api/personal/(?P<id>\\d+)")
def r_personal_quitar(c, p, b, q):
    return quitar_empleado(c, int(p["id"]))


@ruta("POST", "/api/personal/(?P<id>\\d+)/entrada")
def r_entrada(c, p, b, q):
    return marcar_entrada(c, int(p["id"]))


@ruta("POST", "/api/personal/(?P<id>\\d+)/salida")
def r_salida(c, p, b, q):
    return marcar_salida(c, int(p["id"]))


@ruta("GET", "/api/turnos")
def r_turnos(c, p, b, q):
    hoy = hora_actual().strftime("%Y-%m-%d")
    return reporte_turnos(c, q.get("desde") or hoy, q.get("hasta") or hoy)


@ruta("POST", "/api/turnos/(?P<id>\\d+)")
def r_turno_editar(c, p, b, q):
    return editar_turno(c, int(p["id"]), b)


@ruta("DELETE", "/api/turnos/(?P<id>\\d+)")
def r_turno_borrar(c, p, b, q):
    c.execute("delete from turnos where id=?", (int(p["id"]),))
    c.commit()
    return {"ok": True}


@ruta("GET", "/api/version")
def r_version(c, p, b, q):
    return {"version": VERSION}


@ruta("GET", "/api/snapshot")
def r_snapshot(c, p, b, q):
    return snapshot(c)


@ruta("POST", "/api/restaurar")
def r_restaurar(c, p, b, q):
    return restaurar(c, b)


@ruta("POST", "/api/sincronizar")
def r_sincronizar(c, p, b, q):
    if LOCAL:
        raise Err("Operación no permitida")
    return sincronizar(b.get("ops") or [])


@ruta("GET", "/api/clientes")
def r_cliente(c, p, b, q):
    return buscar_cliente(c, q.get("tel")) or {}


@ruta("GET", "/api/impresora")
def r_impresora(c, p, b, q):
    cfg = cfg_get(c)
    return {"puerto": cfg.get("impresora_puerto", ""), "sin_tildes": str(cfg.get("impresora_sin_tildes", "1")) != "0",
            "puertos": impresora.listar_puertos()}


@ruta("POST", "/api/impresora")
def r_impresora_guardar(c, p, b, q):
    puerto = str(b.get("puerto") or "").strip()
    if puerto and not re.fullmatch(r"COM\d{1,3}", puerto.upper()) and os.name == "nt":
        raise Err("El puerto debe verse como COM5")
    c.execute("insert or replace into config(clave,valor) values('impresora_puerto',?)", (puerto.upper() if puerto.upper().startswith("COM") else puerto,))
    c.execute("insert or replace into config(clave,valor) values('impresora_sin_tildes',?)", ("0" if b.get("sin_tildes") is False else "1",))
    c.commit()
    return {"ok": True}


@ruta("POST", "/api/impresora/prueba")
def r_impresora_prueba(c, p, b, q):
    cfg = cfg_get(c)
    return impresora.imprimir(cfg, impresora.texto_prueba(cfg.get("negocio", ""), ANCHO_TICKET), doble_linea=2)


@ruta("POST", "/api/cuentas/(?P<id>\\d+)/reimprimir")
def r_reimprimir(c, p, b, q):
    cid = int(p["id"])
    ronda = c.execute("select max(ronda) from items where cuenta_id=?", (cid,)).fetchone()[0]
    if not ronda:
        raise Err("La cuenta no tiene pedidos")
    texto = texto_comanda(c, cid, int(b.get("ronda") or ronda))
    res = impresora.imprimir(cfg_get(c), texto, doble_linea=2)
    res["comanda"] = texto
    return res


@ruta("POST", "/api/config")
def r_config(c, p, b, q):
    for k in ("negocio", "telefono", "bandeja_domicilio"):
        if k in b:
            c.execute("insert or replace into config(clave,valor) values(?,?)", (k, str(b[k])))
    c.commit()
    return {"ok": True}


class Handler(BaseHTTPRequestHandler):
    server_version = "KatriPOS/1.0"
    protocol_version = "HTTP/1.1"

    def log_message(self, *a):
        pass

    def _send(self, code, body, ctype="application/json; charset=utf-8", extra=None):
        if isinstance(body, str):
            body = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-cache")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _estatico(self, path):
        if path in ("/", "/mesero", "/index.html"):
            path = "/index.html"
        elif path == "/logo.png":
            path = "/logo.png"
        elif path in ("/manifest.webmanifest", "/sw.js"):
            pass
        elif path in ("/py/server.py", "/py/impresora.py"):
            f = os.path.join(BASE, path.split("/")[-1])
            with open(f, "rb") as fh:
                return self._send(200, fh.read(), "text/plain; charset=utf-8")
        rel = path.lstrip("/")
        if rel.startswith("static/"):
            rel = rel[7:]
        full = os.path.abspath(os.path.join(STATIC, rel))
        if not full.startswith(os.path.abspath(STATIC)) or not os.path.isfile(full):
            return self._send(404, json.dumps({"error": "No encontrado"}))
        ctype = {".webmanifest": "application/manifest+json", ".wasm": "application/wasm", ".whl": "application/zip",
                 ".json": "application/json", ".zip": "application/zip"}.get(os.path.splitext(full)[1]) or mimetypes.guess_type(full)[0] or "application/octet-stream"
        if ctype.startswith("text/") or ctype in ("application/javascript",):
            ctype += "; charset=utf-8"
        with open(full, "rb") as f:
            self._send(200, f.read(), ctype)

    def _manejar(self, metodo):
        u = urlparse(self.path)
        path = u.path
        q = {k: v[0] for k, v in parse_qs(u.query).items()}
        if not path.startswith("/api/"):
            if metodo == "GET":
                return self._estatico(path)
            return self._send(404, json.dumps({"error": "No encontrado"}))
        if path == "/api/apagar" and metodo == "POST":
            if self.client_address[0] not in ("127.0.0.1", "::1"):
                return self._send(403, json.dumps({"error": "No permitido"}))
            self._send(200, json.dumps({"ok": True}))
            threading.Timer(0.6, os._exit, [0]).start()
            return
        body = {}
        if metodo in ("POST", "PUT"):
            n = int(self.headers.get("Content-Length") or 0)
            raw = self.rfile.read(n) if n else b""
            try:
                body = json.loads(raw.decode("utf-8")) if raw else {}
            except Exception:
                return self._send(400, json.dumps({"error": "JSON inválido"}))
        if path == "/api/reportes/excel" and metodo == "GET":
            hoy = hora_actual().strftime("%Y-%m-%d")
            d, h = q.get("desde") or hoy, q.get("hasta") or hoy
            try:
                data = excel_bytes(d, h)
            except ImportError:
                return self._send(500, json.dumps({"error": "Falta la librería openpyxl"}))
            return self._send(200, data,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                              {"Content-Disposition": 'attachment; filename="ventas_%s_a_%s.xlsx"' % (d, h)})
        if path == "/api/personal/excel" and metodo == "GET":
            hoy = hora_actual().strftime("%Y-%m-%d")
            d, h = q.get("desde") or hoy, q.get("hasta") or hoy
            try:
                data = excel_personal_bytes(d, h)
            except ImportError:
                return self._send(500, json.dumps({"error": "Falta la librería openpyxl"}))
            return self._send(200, data,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                              {"Content-Disposition": 'attachment; filename="personal_%s_a_%s.xlsx"' % (d, h)})
        estado, res = despachar(metodo, path, q, body)
        return self._send(estado, json.dumps(res, ensure_ascii=False))

    def do_GET(self):
        self._manejar("GET")

    def do_POST(self):
        self._manejar("POST")

    def do_DELETE(self):
        self._manejar("DELETE")


class Servidor(ThreadingHTTPServer):
    daemon_threads = True
    allow_reuse_address = bool(os.environ.get("KATRI_REUSE"))  # solo para pruebas en Linux


def crear_servidor(puerto=PORT):
    init_db()
    return Servidor(("0.0.0.0", puerto), Handler)


if __name__ == "__main__":
    srv = crear_servidor()
    print("Servidor en http://%s:%d" % (ip_local(), PORT))
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass
